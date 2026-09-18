import streamlit as st
import gspread
from google.oauth2.service_account import Credentials
import pandas as pd
from datetime import datetime, date
import pytz
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ==========================================
# 設定
# ==========================================
SPREADSHEET_ID = "1ThtSEj2dnEYSKHIXerjcujo9-6rxmRXEYk2ijS80OlQ"
COMPANY_NAME = "株式会社ハイビックス"
ADMIN_PASSWORD = "3131"
JST = pytz.timezone("Asia/Tokyo")

# ==========================================
# ページ設定
# ==========================================
st.set_page_config(
    page_title=f"{COMPANY_NAME} 管理画面",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ==========================================
# スタイル設定
# ==========================================
st.markdown("""
    <style>
        .stApp {
            background-color: #f0f4f8;
        }
        #MainMenu {visibility: hidden;}
        footer {visibility: hidden;}
        header {visibility: hidden;}
    </style>
""", unsafe_allow_html=True)

# ==========================================
# 日付正規化関数
# ==========================================
def normalize_date(date_str):
    """日付を正規化する（2024/1/15 → 2024/01/15）"""
    try:
        parts = str(date_str).split("/")
        if len(parts) == 3:
            year = int(parts[0])
            month = int(parts[1])
            day = int(parts[2])
            return f"{year:04d}/{month:02d}/{day:02d}"
    except:
        pass
    return date_str

# ==========================================
# スプレッドシート接続
# ==========================================
@st.cache_resource
def get_gspread_client():
    scope = [
        "https://spreadsheets.google.com/feeds",
        "https://www.googleapis.com/auth/drive"
    ]
    creds = Credentials.from_service_account_info(
        st.secrets["gcp_service_account"],
        scopes=scope
    )
    return gspread.Client(auth=creds)

def get_sheet(sheet_name):
    client = get_gspread_client()
    spreadsheet = client.open_by_key(SPREADSHEET_ID)
    return spreadsheet.worksheet(sheet_name)

def get_sheet_data(sheet_name):
    try:
        sheet = get_sheet(sheet_name)
        data = sheet.get_all_records()
        return pd.DataFrame(data)
    except Exception as e:
        st.error(f"読み込みエラー: {e}")
        return pd.DataFrame()

def delete_row(sheet_name, row_index):
    try:
        sheet = get_sheet(sheet_name)
        sheet.delete_rows(row_index + 2)
        return True
    except Exception as e:
        st.error(f"削除エラー: {e}")
        return False

# ==========================================
# ログイン画面
# ==========================================
def show_login():
    st.markdown("""
        <div style="text-align:center; margin-top:100px;">
            <h1>🏭 管理画面ログイン</h1>
            <p style="color:#666;">パスワードを入力してください</p>
        </div>
    """, unsafe_allow_html=True)

    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        password = st.text_input(
            "パスワード",
            type="password",
            placeholder="パスワードを入力",
        )
        login_btn = st.button("🔓 ログイン", use_container_width=True)

        if login_btn:
            if password == ADMIN_PASSWORD:
                st.session_state.logged_in = True
                st.success("ログイン成功！")
                st.rerun()
            else:
                st.error("パスワードが違います")

# ==========================================
# メイン処理
# ==========================================
def main():
    if "logged_in" not in st.session_state:
        st.session_state.logged_in = False

    if not st.session_state.logged_in:
        show_login()
        return

    st.title(f"🏭 {COMPANY_NAME} 管理画面")

    with st.sidebar:
        st.success("✅ ログイン中")
        if st.button("🔒 ログアウト", use_container_width=True):
            st.session_state.logged_in = False
            st.rerun()

    menu = st.sidebar.selectbox(
        "メニューを選択",
        ["🏥 欠勤登録", "🤝 来客登録", "📢 お知らせ登録", "📊 勤怠一覧"]
    )


    # ==========================================
    # 欠勤登録
    # ==========================================
    if menu == "🏥 欠勤登録":
        st.header("🏥 欠勤・遅刻・早退 登録")

        with st.form("absence_form"):
            col1, col2 = st.columns(2)

            with col1:
                name = st.text_input("氏名", placeholder="例：山田 太郎")
                department = st.text_input("部署", placeholder="例：製造部")
                absence_type = st.selectbox(
                    "種別",
                    ["欠勤", "遅刻", "早退", "半休"]
                )

            with col2:
                absence_date = st.date_input("日付", value=datetime.now(JST).date())
                note = st.text_area("備考", placeholder="例：発熱のため", height=100)

            submitted = st.form_submit_button("✅ 登録する", use_container_width=True)

            if submitted:
                if not name or not department:
                    st.error("氏名と部署を入力してください")
                else:
                    try:
                        sheet = get_sheet("欠勤連絡")
                        now = datetime.now(JST).strftime("%H:%M")
                        date_str = absence_date.strftime("%Y/%m/%d")
                        sheet.append_row([
                            date_str,
                            name,
                            department,
                            absence_type,
                            note,
                            now
                        ])
                        st.success(f"✅ {name}さんの{absence_type}を登録しました！")
                        st.cache_resource.clear()
                    except Exception as e:
                        st.error(f"登録エラー: {e}")

        # 登録済みデータ表示・削除
        st.subheader("📋 本日の登録済みデータ")
        df = get_sheet_data("欠勤連絡")

        if not df.empty and "日付" in df.columns:
            today = datetime.now(JST).strftime("%Y/%m/%d")
            df["日付_正規化"] = df["日付"].apply(normalize_date)
            df_today = df[df["日付_正規化"] == today].copy()
            df_today.reset_index(drop=False, inplace=True)

            if not df_today.empty:
                for _, row in df_today.iterrows():
                    col1, col2 = st.columns([4, 1])
                    with col1:
                        st.info(f"👤 {row.get('氏名','')} / {row.get('部署','')} / {row.get('種別','')} / {row.get('備考','')}")
                    with col2:
                        if st.button("🗑️ 削除", key=f"absence_{row['index']}"):
                            if delete_row("欠勤連絡", row["index"]):
                                st.success("削除しました！")
                                st.cache_resource.clear()
                                st.rerun()
            else:
                st.info("本日の登録はありません")
        else:
            st.info("本日の登録はありません")

    # ==========================================
    # 来客登録
    # ==========================================
    elif menu == "🤝 来客登録":
        st.header("🤝 来客情報 登録")

        with st.form("visitor_form"):
            col1, col2 = st.columns(2)

            with col1:
                visit_date = st.date_input("日付", value=datetime.now(JST).date())
                visit_time = st.time_input("来訪時刻", value=datetime.now(JST).time())
                company = st.text_input("会社名", placeholder="例：〇〇株式会社")

            with col2:
                visitor_name = st.text_input("訪問者名", placeholder="例：鈴木 一郎")
                num_people = st.number_input("人数", min_value=1, max_value=50, value=1)
                person_in_charge = st.text_input("担当者", placeholder="例：田中 次郎")

            submitted = st.form_submit_button("✅ 登録する", use_container_width=True)

            if submitted:
                if not company or not visitor_name:
                    st.error("会社名と訪問者名を入力してください")
                else:
                    try:
                        sheet = get_sheet("来客情報")
                        date_str = visit_date.strftime("%Y/%m/%d")
                        time_str = visit_time.strftime("%H:%M")
                        sheet.append_row([
                            date_str,
                            time_str,
                            company,
                            visitor_name,
                            num_people,
                            person_in_charge
                        ])
                        st.success(f"✅ {company}様の来客情報を登録しました！")
                        st.cache_resource.clear()
                    except Exception as e:
                        st.error(f"登録エラー: {e}")

        # 登録済みデータ表示・削除
        st.subheader("📋 本日の登録済みデータ")
        df = get_sheet_data("来客情報")

        if not df.empty and "日付" in df.columns:
            today = datetime.now(JST).strftime("%Y/%m/%d")
            df["日付_正規化"] = df["日付"].apply(normalize_date)
            df_today = df[df["日付_正規化"] == today].copy()
            df_today.reset_index(drop=False, inplace=True)

            if not df_today.empty:
                for _, row in df_today.iterrows():
                    col1, col2 = st.columns([4, 1])
                    with col1:
                        st.info(f"🏢 {row.get('会社名','')} / {row.get('訪問者名','')} / {row.get('来訪時刻','')} / 担当：{row.get('担当者','')}")
                    with col2:
                        if st.button("🗑️ 削除", key=f"visitor_{row['index']}"):
                            if delete_row("来客情報", row["index"]):
                                st.success("削除しました！")
                                st.cache_resource.clear()
                                st.rerun()
            else:
                st.info("本日の登録はありません")
        else:
            st.info("本日の登録はありません")

    # ==========================================
    # お知らせ登録
    # ==========================================
    elif menu == "📢 お知らせ登録":
        st.header("📢 お知らせ 登録")

        with st.form("notice_form"):
            content = st.text_area(
                "お知らせ内容",
                placeholder="例：〇〇のため、△△をお知らせします。",
                height=150
            )
            col1, col2 = st.columns(2)
            with col1:
                start_date = st.date_input("掲載開始日", value=datetime.now(JST).date())
            with col2:
                end_date = st.date_input("掲載終了日", value=datetime.now(JST).date())

            submitted = st.form_submit_button("✅ 登録する", use_container_width=True)

            if submitted:
                if not content:
                    st.error("お知らせ内容を入力してください")
                elif start_date > end_date:
                    st.error("掲載終了日は開始日以降にしてください")
                else:
                    try:
                        sheet = get_sheet("お知らせ")
                        start_str = start_date.strftime("%Y/%m/%d")
                        end_str = end_date.strftime("%Y/%m/%d")
                        sheet.append_row([content, start_str, end_str])
                        st.success("✅ お知らせを登録しました！")
                        st.cache_resource.clear()
                    except Exception as e:
                        st.error(f"登録エラー: {e}")

        # 登録済みデータ表示・削除
        st.subheader("📋 登録済みお知らせ一覧")
        df = get_sheet_data("お知らせ")

        if not df.empty:
            df_notice = df.copy()
            df_notice.reset_index(drop=False, inplace=True)
            for _, row in df_notice.iterrows():
                col1, col2 = st.columns([4, 1])
                with col1:
                    st.info(f"📌 {row.get('お知らせ内容','')} / {row.get('掲載開始日','')} ～ {row.get('掲載終了日','')}")
                with col2:
                    if st.button("🗑️ 削除", key=f"notice_{row['index']}"):
                        if delete_row("お知らせ", row["index"]):
                            st.success("削除しました！")
                            st.cache_resource.clear()
                            st.rerun()
        else:
            st.info("登録済みのお知らせはありません")

    # ==========================================
    # 勤怠一覧
    # ==========================================
    elif menu == "📊 勤怠一覧":
        st.header("📊 勤怠一覧 Excelダウンロード")

        # 期間指定
        col1, col2 = st.columns(2)
        with col1:
            start_date = st.date_input(
                "開始日",
                value=datetime.now(JST).date().replace(day=1)
            )
        with col2:
            end_date = st.date_input(
                "終了日",
                value=datetime.now(JST).date()
            )

        if start_date > end_date:
            st.error("終了日は開始日以降にしてください")
            return

        # フィルター
        df = get_sheet_data("欠勤連絡")

        if df.empty:
            st.info("データがありません")
        else:
            df["日付_正規化"] = df["日付"].apply(normalize_date)
            start_str = start_date.strftime("%Y/%m/%d")
            end_str = end_date.strftime("%Y/%m/%d")
            df_filtered = df[
                (df["日付_正規化"] >= start_str) &
                (df["日付_正規化"] <= end_str)
            ].copy()

            # 部署フィルター
            departments = ["すべて"] + sorted(df_filtered["部署"].dropna().unique().tolist())
            selected_dept = st.selectbox("部署で絞り込み", departments)
            if selected_dept != "すべて":
                df_filtered = df_filtered[df_filtered["部署"] == selected_dept]

            # 名前フィルター
            names = ["すべて"] + sorted(df_filtered["氏名"].dropna().unique().tolist())
            selected_name = st.selectbox("氏名で絞り込み", names)
            if selected_name != "すべて":
                df_filtered = df_filtered[df_filtered["氏名"] == selected_name]

            # 表示列
            display_cols = ["日付", "氏名", "部署", "種別", "開始時刻", "終了時刻", "備考", "登録時刻"]
            df_display = df_filtered[[col for col in display_cols if col in df_filtered.columns]].copy()
            df_display = df_display.sort_values("日付").reset_index(drop=True)

            st.subheader(f"📋 該当件数：{len(df_display)}件")
            st.dataframe(df_display, use_container_width=True)

            # Excelファイル作成
            def create_excel(df):
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "勤怠一覧"

                # タイトル
                ws.merge_cells("A1:H1")
                title_cell = ws["A1"]
                title_cell.value = f"勤怠一覧　{start_date.strftime('%Y/%m/%d')} ～ {end_date.strftime('%Y/%m/%d')}"
                title_cell.font = Font(bold=True, size=14, color="FFFFFF")
                title_cell.fill = PatternFill("solid", fgColor="1565C0")
                title_cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[1].height = 30

                # ヘッダー
                headers = list(df.columns)
                header_fill = PatternFill("solid", fgColor="0288D1")
                thin = Side(style="thin", color="CCCCCC")
                border = Border(left=thin, right=thin, top=thin, bottom=thin)

                for col_idx, header in enumerate(headers, 1):
                    cell = ws.cell(row=2, column=col_idx, value=header)
                    cell.font = Font(bold=True, color="FFFFFF", size=11)
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = border
                ws.row_dimensions[2].height = 22

                # データ
                for row_idx, row in df.iterrows():
                    for col_idx, value in enumerate(row.values, 1):
                        cell = ws.cell(row=row_idx + 3, column=col_idx, value=value)
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.border = border
                        if row_idx % 2 == 0:
                            cell.fill = PatternFill("solid", fgColor="E3F2FD")
                    ws.row_dimensions[row_idx + 3].height = 20

                # 列幅
                column_widths = {
                    "日付": 14,
                    "氏名": 14,
                    "部署": 14,
                    "種別": 12,
                    "開始時刻": 12,
                    "終了時刻": 12,
                    "備考": 24,
                    "登録時刻": 12,
                }
                for col_idx, header in enumerate(headers, 1):
                    ws.column_dimensions[
                        openpyxl.utils.get_column_letter(col_idx)
                    ].width = column_widths.get(header, 14)

                output = BytesIO()
                wb.save(output)
                output.seek(0)
                return output

            # ダウンロードボタン
            if not df_display.empty:
                excel_data = create_excel(df_display)
                file_name = f"勤怠一覧_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"
                st.download_button(
                    label="📥 Excelダウンロード",
                    data=excel_data,
                    file_name=file_name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
            else:
                st.info("該当するデータがありません")

if __name__ == "__main__":
    main()
